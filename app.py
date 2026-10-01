import io
import json
import os
import re
from copy import deepcopy
from datetime import datetime, timedelta, timezone

import pandas as pd
import streamlit as st
from docx import Document
from groq import Groq

MODEL = "openai/gpt-oss-20b"
PKT = timezone(timedelta(hours=5))  # Pakistan time (no daylight saving)
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

DATE_RE = re.compile(r"(\bDate\s*:\s*)(\d{1,2})([/.\-])(\d{1,2})\3(\d{2,4})", re.I)
DAY_RE = re.compile(r"(\bDay\s*:\s*)([A-Za-z]+)", re.I)
STEP_RE = re.compile(r"(\bStep\s*:\s*)(.+?)\s*$", re.I)


# -----------------------------
# Document helpers
# -----------------------------
def iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        seen = set()
        for row in t.rows:
            for cell in row.cells:
                if cell._tc in seen:  # merged cells repeat
                    continue
                seen.add(cell._tc)
                for p in cell.paragraphs:
                    yield p
    for s in doc.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                yield p


def replace_span(par, start, end, new):
    """Replace characters [start:end] of a paragraph, keeping run formatting."""
    pos, first = 0, True
    for r in par.runs:
        rs, re_ = pos, pos + len(r.text)
        pos = re_
        if re_ <= start or rs >= end:
            continue
        a, b = max(start, rs) - rs, min(end, re_) - rs
        r.text = r.text[:a] + (new if first else "") + r.text[b:]
        first = False


def format_date_like(match, d):
    """Write the new date in the same style as the old one (1/10/26 -> 2/10/26)."""
    td, sep, tm, ty = match.group(2), match.group(3), match.group(4), match.group(5)
    day = f"{d.day:02d}" if td.startswith("0") else str(d.day)
    mon = f"{d.month:02d}" if tm.startswith("0") else str(d.month)
    year = str(d.year) if len(ty) == 4 else f"{d.year % 100:02d}"
    return f"{day}{sep}{mon}{sep}{year}"


def update_date_and_day(doc, new_date, step=None):
    changes = []
    for p in iter_paragraphs(doc):
        m = DATE_RE.search(p.text)
        if m:
            old = p.text[m.start(2):m.end(5)]
            new = format_date_like(m, new_date)
            replace_span(p, m.start(2), m.end(5), new)
            changes.append(f"Date: {old} → {new}")

        m = DAY_RE.search(p.text)
        if m:
            old = m.group(2)
            name = new_date.strftime("%A")
            name = name.upper() if old.isupper() else name
            replace_span(p, m.start(2), m.end(2), name)
            changes.append(f"Day: {old} → {name}")

        if step is not None:
            m = STEP_RE.search(p.text)
            if m and m.group(2) != step.strip():
                old = m.group(2)
                replace_span(p, m.start(2), m.end(2), step.strip())
                changes.append(f"Step: {old} → {step.strip()}")
    return changes


def detect_step(doc):
    for p in iter_paragraphs(doc):
        m = STEP_RE.search(p.text)
        if m:
            return m.group(2)
    return None


def find_diary_table(doc):
    """Find the table that has 'Subject' and 'Homework' header cells."""
    for t in doc.tables:
        for ri, row in enumerate(t.rows):
            labels = {}
            for ci, cell in enumerate(row.cells):
                labels.setdefault(cell.text.strip().lower(), ci)
            if "subject" in labels and "homework" in labels:
                return t, ri, labels["subject"], labels["homework"]
    return None


def read_rows(t, hi, sc, hc):
    return [
        {"Subject": r.cells[sc].text.strip(), "Homework": r.cells[hc].text.strip()}
        for r in t.rows[hi + 1:]
    ]


def find_ref_run(t, hi, sc, hc):
    for r in t.rows[hi + 1:]:
        for ci in (sc, hc):
            for p in r.cells[ci].paragraphs:
                if p.runs and p.text.strip():
                    return p.runs[0]
    return None


def set_cell_text(cell, text, ref=None):
    paras = cell.paragraphs
    first = paras[0]
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    if first.runs:
        first.runs[0].text = text
        for r in first.runs[1:]:
            r._r.getparent().remove(r._r)
    elif text:
        run = first.add_run(text)
        if ref is not None and ref._r.rPr is not None:  # copy font/bold from a filled cell
            run._r.insert(0, deepcopy(ref._r.rPr))


def write_rows(t, hi, sc, hc, rows):
    ref = find_ref_run(t, hi, sc, hc)
    while len(t.rows) - hi - 1 < len(rows):  # add table rows if more are needed
        last = t.rows[-1]._tr
        last.addnext(deepcopy(last))
    for i, row in enumerate(t.rows[hi + 1:]):
        d = rows[i] if i < len(rows) else {"Subject": "", "Homework": ""}
        set_cell_text(row.cells[sc], d["Subject"], ref)
        set_cell_text(row.cells[hc], d["Homework"], ref)


def next_school_day(today):
    d = today + timedelta(days=1)
    while d.weekday() >= 5:  # skip Saturday/Sunday
        d += timedelta(days=1)
    return d


# -----------------------------
# Optional AI fill
# -----------------------------
def get_client():
    key = os.getenv("GROQ_API_KEY")
    if not key:
        try:
            key = st.secrets["GROQ_API_KEY"]
        except Exception:
            key = None
    if not key:
        raise RuntimeError("GROQ_API_KEY is not configured in secrets.")
    return Groq(api_key=key)


def ai_fill(client, message, current_rows, clear_unmentioned):
    rule = (
        "For subjects NOT mentioned in the message, set Homework to an empty string."
        if clear_unmentioned
        else "For subjects NOT mentioned in the message, keep their current Homework unchanged."
    )
    system = (
        "You update a school homework diary table. You get the current rows and a message "
        "describing the new homework. Return ONLY JSON: "
        '{"rows": [{"Subject": "...", "Homework": "..."}]}. '
        "Keep existing subject names, spelling and order. " + rule + " "
        "Add subjects that are in the message but not in the table at the end. "
        "Never invent homework. Keep the wording short, as given in the message."
    )
    user = f"CURRENT ROWS:\n{json.dumps(current_rows)}\n\nMESSAGE:\n{message}"
    resp = client.chat.completions.create(
        model=MODEL,
        temperature=0.1,
        response_format={"type": "json_object"},
        messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
    )
    content = re.sub(r"^```(?:json)?\s*|\s*```$", "", resp.choices[0].message.content.strip())
    rows = json.loads(content).get("rows", [])
    if not rows:
        raise RuntimeError("AI returned no rows. Try rewording the message.")
    return [{"Subject": str(r.get("Subject", "")), "Homework": str(r.get("Homework", ""))} for r in rows]


# -----------------------------
# UI
# -----------------------------
st.set_page_config(page_title="Daily Diary Updater", page_icon="📒", layout="wide")
st.title("📒 Daily Diary Updater")
st.caption("Upload the diary → date & day update automatically → edit homework → download.")

uploaded = st.file_uploader("Upload the diary (.docx)", type=["docx"])
if not uploaded:
    st.info("Upload the diary file to begin.")
    st.stop()

data = uploaded.getvalue()
fkey = f"{uploaded.name}-{len(data)}"
if st.session_state.get("fkey") != fkey:
    st.session_state.fkey = fkey
    st.session_state.rows = None
    st.session_state.editor_v = 0
    st.session_state.output = None

doc = Document(io.BytesIO(data))
found = find_diary_table(doc)

if st.session_state.rows is None:
    st.session_state.rows = read_rows(*found) if found else []

# ---- Date ----
today = datetime.now(PKT).date()
st.subheader("1. Date & day")
mode = st.radio("Diary date", ["Today", "Next school day (Mon–Fri)", "Pick a date"], horizontal=True)
if mode == "Today":
    target = today
elif mode.startswith("Next"):
    target = next_school_day(today)
else:
    target = st.date_input("Choose date", value=today)

current_step = detect_step(doc)
step = None
if current_step is not None:
    step = st.text_input("Step / class", value=current_step, key=f"step_{fkey}")

preview = update_date_and_day(Document(io.BytesIO(data)), target, step)
if preview:
    st.info("  |  ".join(preview))
else:
    st.warning("Couldn't find 'Date:' or 'Day:' in this file.")

# ---- Homework table ----
st.subheader("2. Subjects & homework")
if not found:
    st.warning("No Subject/Homework table found — only the date and day will be updated.")
    rows_list = []
else:
    df = pd.DataFrame(st.session_state.rows or [{"Subject": "", "Homework": ""}],
                      columns=["Subject", "Homework"])
    edited = st.data_editor(
        df, num_rows="dynamic", hide_index=True,
        key=f"ed_{fkey}_{st.session_state.editor_v}",
    ).fillna("")
    rows_list = [
        {"Subject": str(a).strip(), "Homework": str(b).strip()}
        for a, b in zip(edited["Subject"], edited["Homework"])
    ]
    while rows_list and not rows_list[-1]["Subject"] and not rows_list[-1]["Homework"]:
        rows_list.pop()

    with st.expander("✨ Optional: paste the homework message and let AI fill the table"):
        msg = st.text_area("Homework message", placeholder="Maths ex 6 A, English essay on my school, Science unit 4 Q 1,2")
        clear = st.checkbox("Clear homework for subjects not mentioned", value=True)
        if st.button("Fill table with AI"):
            if not msg.strip():
                st.warning("Paste the homework message first.")
            else:
                try:
                    st.session_state.rows = ai_fill(get_client(), msg, rows_list, clear)
                    st.session_state.editor_v += 1
                    st.rerun()
                except Exception as e:
                    st.error(f"Error: {e}")

# ---- Generate ----
st.subheader("3. Download")
if st.button("✅ Generate updated diary", type="primary"):
    out_doc = Document(io.BytesIO(data))
    update_date_and_day(out_doc, target, step)
    table = find_diary_table(out_doc)
    if table:
        write_rows(*table, rows_list)
    buf = io.BytesIO()
    out_doc.save(buf)
    st.session_state.output = buf.getvalue()
    st.session_state.out_name = f"diary_{target:%d-%m-%Y}.docx"

if st.session_state.get("output"):
    st.success("Diary updated.")
    st.download_button("⬇️ Download updated diary", data=st.session_state.output,
                       file_name=st.session_state.out_name, mime=DOCX_MIME)

st.divider()
st.caption("Files are processed in memory and not stored. AI fill sends only the homework text to Groq.")
